from sqlalchemy.orm import Session
from sqlalchemy import func, desc, and_
from fastapi import HTTPException
from decimal import Decimal
from datetime import date, datetime, timedelta
from typing import Optional, List
import httpx
import json
import asyncio

from app.models.models import (
    Invoice, InvoiceItem, Customer, CustomerAddress, Warehouse,
    Product, Purchase, PurchaseItem, Vendor,
    EInvoiceLog, EWayBillLog, EInvoiceStatus, EWayBillStatus,
    Transporter, Vehicle, DocumentType, GSTType
)
from app.schemas.gst import (
    EInvoiceGenerateRequest, EInvoiceCancelRequest,
    EWayBillGenerateRequest, GSTR2BImport,
    TransporterCreate, TransporterUpdate, VehicleCreate
)
from app.core.config import settings
from app.utils.helpers import paginate, get_financial_year
from app.utils.irp_validation import validate_einvoice, normalize_uqc, compute_round_off
from app.services.audit import audit
from app.services.company_service import CompanyService
from app.services.config_service import ConfigService
from app.services.gst_master_service import GstMasterService


class TransporterService:

    @staticmethod
    def list_transporters(db: Session) -> List[dict]:
        rows = db.query(Transporter).filter(Transporter.is_active == True).order_by(Transporter.name).all()
        return [{"id": t.id, "name": t.name, "gstin": t.gstin,
                 "contact_person": t.contact_person, "phone": t.phone,
                 "email": t.email, "is_active": t.is_active, "created_at": t.created_at}
                for t in rows]

    @staticmethod
    def create_transporter(db: Session, payload: TransporterCreate, user_id: int) -> dict:
        t = Transporter(**payload.dict(), created_by=user_id)
        db.add(t)
        db.commit()
        db.refresh(t)
        audit(db, user_id, "create", "gst",
              f"Created transporter {t.name}" + (f" (GSTIN {t.gstin})" if t.gstin else ""),
              record_type="transporter", record_id=t.id)
        return {"id": t.id, "name": t.name, "gstin": t.gstin,
                "contact_person": t.contact_person, "phone": t.phone,
                "email": t.email, "is_active": t.is_active, "created_at": t.created_at}

    @staticmethod
    def update_transporter(db: Session, transporter_id: int, payload: TransporterUpdate, user_id: int) -> dict:
        t = db.query(Transporter).filter(Transporter.id == transporter_id, Transporter.is_active == True).first()
        if not t:
            raise ValueError("Transporter not found")
        for k, v in payload.dict(exclude_unset=True).items():
            setattr(t, k, v)
        db.commit()
        db.refresh(t)
        audit(db, user_id, "update", "gst", f"Updated transporter {t.name}",
              record_type="transporter", record_id=t.id)
        return {"id": t.id, "name": t.name, "gstin": t.gstin,
                "contact_person": t.contact_person, "phone": t.phone,
                "email": t.email, "is_active": t.is_active}

    @staticmethod
    def delete_transporter(db: Session, transporter_id: int, user_id: int) -> dict:
        t = db.query(Transporter).filter(Transporter.id == transporter_id, Transporter.is_active == True).first()
        if not t:
            raise ValueError("Transporter not found")
        t.is_active = False
        db.commit()
        audit(db, user_id, "delete", "gst", f"Deactivated transporter {t.name}",
              record_type="transporter", record_id=t.id)
        return {"id": t.id, "name": t.name, "is_active": False}

    @staticmethod
    def list_vehicles(db: Session, transporter_id: Optional[int] = None) -> List[dict]:
        q = db.query(Vehicle).filter(Vehicle.is_active == True)
        if transporter_id:
            q = q.filter(Vehicle.transporter_id == transporter_id)
        rows = q.order_by(Vehicle.vehicle_number).all()
        result = []
        for v in rows:
            trans = db.query(Transporter).filter(Transporter.id == v.transporter_id).first() if v.transporter_id else None
            result.append({
                "id": v.id, "vehicle_number": v.vehicle_number,
                "vehicle_type": v.vehicle_type, "owner_name": v.owner_name,
                "transporter_id": v.transporter_id,
                "transporter_name": trans.name if trans else None,
                "is_active": v.is_active, "created_at": v.created_at,
            })
        return result

    @staticmethod
    def create_vehicle(db: Session, payload: VehicleCreate, user_id: int) -> dict:
        v = Vehicle(**payload.dict(), created_by=user_id)
        db.add(v)
        db.commit()
        db.refresh(v)
        audit(db, user_id, "create", "gst",
              f"Created vehicle {v.vehicle_number}",
              record_type="vehicle", record_id=v.id)
        trans = db.query(Transporter).filter(Transporter.id == v.transporter_id).first() if v.transporter_id else None
        return {
            "id": v.id, "vehicle_number": v.vehicle_number,
            "vehicle_type": v.vehicle_type, "owner_name": v.owner_name,
            "transporter_id": v.transporter_id,
            "transporter_name": trans.name if trans else None,
            "is_active": v.is_active, "created_at": v.created_at,
        }

    @staticmethod
    def update_vehicle(db: Session, vehicle_id: int, payload, user_id: int) -> dict:
        v = db.query(Vehicle).filter(Vehicle.id == vehicle_id, Vehicle.is_active == True).first()
        if not v:
            raise ValueError("Vehicle not found")
        data = payload.dict(exclude_unset=True)
        for k, val in data.items():
            setattr(v, k, val)
        db.commit()
        db.refresh(v)
        audit(db, user_id, "update", "gst",
              f"Updated vehicle {v.vehicle_number}",
              record_type="vehicle", record_id=v.id)
        trans = db.query(Transporter).filter(Transporter.id == v.transporter_id).first() if v.transporter_id else None
        return {
            "id": v.id, "vehicle_number": v.vehicle_number,
            "vehicle_type": v.vehicle_type, "owner_name": v.owner_name,
            "transporter_id": v.transporter_id,
            "transporter_name": trans.name if trans else None,
            "is_active": v.is_active, "created_at": v.created_at,
        }

    @staticmethod
    def delete_vehicle(db: Session, vehicle_id: int, user_id: int) -> dict:
        v = db.query(Vehicle).filter(Vehicle.id == vehicle_id, Vehicle.is_active == True).first()
        if not v:
            raise ValueError("Vehicle not found")
        v.is_active = False
        db.commit()
        audit(db, user_id, "delete", "gst",
              f"Deactivated vehicle {v.vehicle_number}",
              record_type="vehicle", record_id=v.id)
        return {"id": v.id, "vehicle_number": v.vehicle_number, "is_active": False}


class EInvoiceService:

    @staticmethod
    def _should_mock() -> bool:
        """Mock when sandbox is enabled and no real auth token is configured.

        A token is considered unconfigured when it is blank or looks like a
        placeholder copied from .env.example (e.g. "your-cleartax-token-here").
        """
        token = (settings.CLEARTAX_AUTH_TOKEN or "").strip()
        token_is_placeholder = not token or "your-" in token.lower() or "token-here" in token.lower()
        return settings.CLEARTAX_SANDBOX and token_is_placeholder

    @staticmethod
    def _derive_suptyp(invoice: Invoice, customer: Optional[Customer]) -> str:
        with_tax = float(invoice.total_igst or 0) > 0
        if customer and getattr(customer, "is_export", False):
            return "EXPWP" if with_tax else "EXPWOP"
        if customer and getattr(customer, "is_sez", False):
            return "SEZWP" if with_tax else "SEZWOP"
        return "B2B"

    @staticmethod
    def _parse_ack_date(ack_date) -> Optional[datetime]:
        if not ack_date:
            return None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(str(ack_date), fmt)
            except (ValueError, TypeError):
                continue
        return None

    @staticmethod
    def _build_payload(db: Session, invoice: Invoice) -> dict:
        company = CompanyService.identity(db)
        customer = db.query(Customer).filter(Customer.id == invoice.customer_id).first()
        bill_addr = db.query(CustomerAddress).filter(CustomerAddress.id == invoice.billing_address_id).first()
        ship_addr = db.query(CustomerAddress).filter(CustomerAddress.id == invoice.shipping_address_id).first()

        uqc_map = GstMasterService.uqc_map(db, invoice.invoice_date)
        item_list = []
        for i, item in enumerate(invoice.items, 1):
            product = db.query(Product).filter(Product.id == item.product_id).first()
            item_list.append({
                "SlNo": str(i),
                "PrdDesc": product.part_name if product else f"Product {item.product_id}",
                "IsServc": "N",
                "HsnCd": item.hsn_code,
                "Qty": float(item.quantity),
                "Unit": normalize_uqc(product.unit_of_measure if product else None, uqc_map),
                "UnitPrice": float(item.unit_price),
                "TotAmt": round(float(item.quantity) * float(item.unit_price), 2),
                "Discount": float(item.discount_amount),
                "AssAmt": float(item.taxable_amount),
                "GstRt": float(item.gst_percent),
                "IgstAmt": float(item.igst_amount),
                "CgstAmt": float(item.cgst_amount),
                "SgstAmt": float(item.sgst_amount),
                "TotItemVal": float(item.line_total),
            })

        return {
            "Version": "1.1",
            "TranDtls": {
                "TaxSch": "GST",
                "SupTyp": EInvoiceService._derive_suptyp(invoice, customer),
                "RegRev": "N",
                "EcmGstin": None,
                "IgstOnIntra": "N",
            },
            "DocDtls": {
                "Typ": "INV",
                "No": invoice.invoice_number,
                "Dt": invoice.invoice_date.strftime("%d/%m/%Y"),
            },
            "SellerDtls": {
                "Gstin": company.gstin,
                "LglNm": company.name,
                "TrdNm": company.name,
                "Addr1": company.address,
                "Loc": company.city,
                "Pin": int(company.pincode),
                "Stcd": str(company.state_code),
                "Ph": company.phone,
                "Em": company.email,
            },
            "BuyerDtls": {
                "Gstin": customer.gstin or "URP",
                "LglNm": customer.legal_name or customer.trade_name,
                "TrdNm": customer.trade_name,
                "Pos": str(invoice.place_of_supply or (ship_addr.state_code if ship_addr else company.state_code)),
                "Addr1": bill_addr.address_line1 if bill_addr else "",
                "Loc": bill_addr.city if bill_addr else "",
                "Pin": int(bill_addr.pincode) if bill_addr and bill_addr.pincode else 0,
                "Stcd": str(bill_addr.state_code if bill_addr and bill_addr.state_code else company.state_code),
                "Ph": customer.phone or "",
                "Em": customer.email or "",
            },
            "DispDtls": {
                "Nm": company.name,
                "Addr1": company.address,
                "Loc": company.city,
                "Pin": int(company.pincode),
                "Stcd": str(company.state_code),
            },
            "ShipDtls": {
                "Gstin": customer.gstin or "URP",
                "LglNm": customer.trade_name,
                "TrdNm": customer.trade_name,
                "Addr1": ship_addr.address_line1 if ship_addr else "",
                "Loc": ship_addr.city if ship_addr else "",
                "Pin": int(ship_addr.pincode) if ship_addr and ship_addr.pincode else 0,
                "Stcd": str(ship_addr.state_code if ship_addr and ship_addr.state_code else company.state_code),
            },
            "ItemList": item_list,
            "ValDtls": {
                "AssVal": float(invoice.taxable_amount),
                "CgstVal": float(invoice.total_cgst),
                "SgstVal": float(invoice.total_sgst),
                "IgstVal": float(invoice.total_igst),
                "TotInvVal": float(invoice.total_amount),
                "RndOffAmt": compute_round_off(invoice),
                "Discount": float(invoice.item_discount + invoice.invoice_discount),
            },
        }

    @staticmethod
    async def generate(db: Session, invoice_id: int, user_id: int) -> dict:
        invoice = db.query(Invoice).filter(Invoice.id == invoice_id).first()
        if not invoice:
            raise HTTPException(status_code=404, detail="Invoice not found")
        if invoice.is_cancelled:
            raise HTTPException(status_code=400, detail="Cannot generate IRN for cancelled invoice")
        if invoice.irn:
            raise HTTPException(status_code=400, detail=f"IRN already generated: {invoice.irn}")
        if invoice.document_type not in (DocumentType.b2b_invoice,):
            raise HTTPException(status_code=400, detail="E-Invoice only applicable for B2B invoices")

        # Pre-flight IRP validation: fail fast before spending an API call.
        customer = db.query(Customer).filter(Customer.id == invoice.customer_id).first()
        bill_addr = db.query(CustomerAddress).filter(
            CustomerAddress.id == invoice.billing_address_id).first()
        ship_addr = db.query(CustomerAddress).filter(
            CustomerAddress.id == invoice.shipping_address_id).first()
        company = CompanyService.identity(db)
        validate_einvoice(
            invoice, customer, bill_addr, ship_addr,
            company.gstin, company.state_code,
            min_hsn_digits=ConfigService.get_int(
                db, "gst.hsn_min_digits",
                as_of=invoice.invoice_date, default=settings.HSN_MIN_DIGITS,
            ),
            valid_gst_rates=GstMasterService.valid_gst_rates(db, invoice.invoice_date) or None,
            valid_state_codes=GstMasterService.valid_state_codes(db, invoice.invoice_date) or None,
        )

        payload = EInvoiceService._build_payload(db, invoice)

        # Log the attempt
        log = EInvoiceLog(
            invoice_id=invoice_id,
            status=EInvoiceStatus.pending,
            request_payload=json.dumps(payload),
            created_by=user_id,
        )
        db.add(log)
        db.flush()

        # Try Cleartax API with 3 retries
        irn = qr_code = ack_number = ack_date = signed_invoice = None
        error_message = None
        success = False

        attempts = 0 if EInvoiceService._should_mock() else 3
        for attempt in range(attempts):
            try:
                async with httpx.AsyncClient(timeout=30) as client:
                    resp = await client.post(
                        f"{settings.CLEARTAX_API_URL}/v2/eInvoice/generate",
                        json=payload,
                        headers={
                            "x-cleartax-auth-token": settings.CLEARTAX_AUTH_TOKEN,
                            "Content-Type": "application/json",
                        }
                    )
                    if resp.status_code == 200:
                        data = resp.json()
                        irn = data.get("Irn")
                        qr_code = data.get("SignedQRCode")
                        signed_invoice = data.get("SignedInvoice")
                        ack_number = data.get("AckNo")
                        ack_date = data.get("AckDt")
                        success = True
                        break
                    else:
                        error_message = resp.text
            except Exception as e:
                error_message = str(e)
                if attempt < 2:
                    await asyncio.sleep(2 ** attempt)

        if success and irn:
            invoice.irn = irn
            invoice.irn_status = EInvoiceStatus.generated
            invoice.qr_code = qr_code
            invoice.signed_invoice = signed_invoice
            invoice.irn_ack_number = str(ack_number) if ack_number else None
            invoice.irn_ack_date = EInvoiceService._parse_ack_date(ack_date)
            invoice.irn_generated_at = datetime.now()
            invoice.irn_error = None
            log.status = EInvoiceStatus.generated
            log.irn = irn
            log.ack_number = invoice.irn_ack_number
            log.ack_date = invoice.irn_ack_date
            log.response_payload = json.dumps({
                "irn": irn, "ack_number": ack_number, "ack_date": ack_date
            })
        else:
            # Use sandbox/mock IRN only when no real token is configured
            if EInvoiceService._should_mock():
                mock_irn = f"MOCK-IRN-{invoice.invoice_number}-{invoice_id}"
                ack_number = f"MOCK-ACK-{invoice_id}"
                invoice.irn = mock_irn
                invoice.irn_status = EInvoiceStatus.generated
                invoice.qr_code = f"MOCK-QR-{invoice_id}"
                invoice.irn_ack_number = ack_number
                invoice.irn_generated_at = datetime.now()
                invoice.irn_error = None
                log.status = EInvoiceStatus.generated
                log.irn = mock_irn
                log.ack_number = ack_number
                irn = mock_irn
                success = True
                error_message = None
            else:
                invoice.irn_status = EInvoiceStatus.failed
                invoice.irn_error = error_message
                log.status = EInvoiceStatus.failed
                log.error_message = error_message

        db.commit()

        audit(db, user_id, "create", "gst",
              f"Generated IRN for invoice {invoice.invoice_number}: {irn}" if success
              else f"IRN generation failed for invoice {invoice.invoice_number}: {error_message}",
              record_type="einvoice", record_id=invoice_id)

        return {
            "invoice_id": invoice_id,
            "invoice_number": invoice.invoice_number,
            "irn": irn,
            "irn_status": invoice.irn_status,
            "qr_code": invoice.qr_code,
            "signed_invoice": invoice.signed_invoice,
            "ack_number": invoice.irn_ack_number,
            "ack_date": ack_date,
            "error_message": error_message if not success else None,
        }

    @staticmethod
    async def cancel_irn(db: Session, invoice_id: int, reason: str, user_id: int,
                         remark: Optional[str] = None) -> dict:
        invoice = db.query(Invoice).filter(Invoice.id == invoice_id).first()
        if not invoice:
            raise HTTPException(status_code=404, detail="Invoice not found")
        if not invoice.irn:
            raise HTTPException(status_code=400, detail="No IRN found for this invoice")
        if invoice.irn_status == EInvoiceStatus.cancelled:
            raise HTTPException(status_code=400, detail="IRN already cancelled")
        cancel_window_hours = ConfigService.get_int(
            db, "gst.irn_cancel_window_hours",
            as_of=invoice.irn_generated_at.date() if invoice.irn_generated_at else None,
            default=24,
        )
        if invoice.irn_generated_at and (datetime.now() - invoice.irn_generated_at) > timedelta(hours=cancel_window_hours):
            raise HTTPException(
                status_code=400,
                detail=f"E-Invoice can only be cancelled within {cancel_window_hours} hours of IRN generation"
            )

        # CnlRsn per NIC: 1=Duplicate, 2=Data entry mistake, 3=Order cancelled, 4=Others.
        reason_str = (reason or "").strip()
        reason_code = reason_str if reason_str in ("1", "2", "3", "4") else "2"
        cnl_remark = (remark or reason_str or "Cancelled")[:100]
        cancel_payload = {"Irn": invoice.irn, "CnlRsn": reason_code, "CnlRem": cnl_remark}

        success = False
        error_message = None
        response_data = None

        if EInvoiceService._should_mock():
            success = True
            response_data = {"mock": True, "CancelDate": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
        else:
            try:
                async with httpx.AsyncClient(timeout=30) as client:
                    resp = await client.post(
                        f"{settings.CLEARTAX_API_URL}/v2/eInvoice/cancel",
                        json=cancel_payload,
                        headers={"x-cleartax-auth-token": settings.CLEARTAX_AUTH_TOKEN},
                    )
                    if resp.status_code == 200:
                        success = True
                        try:
                            response_data = resp.json()
                        except Exception:
                            response_data = {"raw": resp.text}
                    else:
                        error_message = resp.text
            except Exception as e:
                error_message = str(e) or "Network error connecting to Cleartax API"

        log = EInvoiceLog(
            invoice_id=invoice_id,
            irn=invoice.irn,
            status=EInvoiceStatus.cancelled if success else EInvoiceStatus.failed,
            request_payload=json.dumps(cancel_payload),
            response_payload=json.dumps(response_data) if response_data is not None else None,
            error_message=error_message,
            created_by=user_id,
        )
        db.add(log)
        if success:
            invoice.irn_status = EInvoiceStatus.cancelled
        db.commit()

        audit(db, user_id, "cancel", "gst",
              f"Cancelled IRN for invoice {invoice.invoice_number}. Reason: {reason}" if success
              else f"IRN cancellation failed for invoice {invoice.invoice_number}: {error_message}",
              record_type="einvoice", record_id=invoice_id)

        return {
            "invoice_id": invoice_id,
            "invoice_number": invoice.invoice_number,
            "irn": invoice.irn,
            "irn_status": invoice.irn_status,
            "error_message": error_message if not success else None,
        }

    @staticmethod
    def list_einvoice_logs(db: Session, page: int = 1, page_size: int = 20) -> dict:
        q = db.query(EInvoiceLog).order_by(desc(EInvoiceLog.created_at))
        result = paginate(q, page, page_size)
        items = []
        for log in result["items"]:
            inv = db.query(Invoice).filter(Invoice.id == log.invoice_id).first()
            items.append({
                "id": log.id,
                "invoice_id": log.invoice_id,
                "invoice_number": inv.invoice_number if inv else None,
                "status": log.status,
                "irn": log.irn,
                "error_message": log.error_message,
                "created_at": log.created_at,
            })
        result["items"] = items
        return result


class EWayBillService:

    @staticmethod
    def _calc_distance(from_pincode: str, to_pincode: str) -> int:
        """Rough offline estimate from pincode prefix difference.

        Used ONLY as a display/validity fallback in mock mode or when the
        IRP response omits the distance. For real submissions we send
        transDistance=0 so NIC auto-calculates the actual pin-to-pin road
        distance; this estimate is never claimed as the official distance.
        """
        try:
            diff = abs(int(from_pincode[:3]) - int(to_pincode[:3]))
            return max(10, diff * 15)
        except Exception:
            return 100

    @staticmethod
    def _extract_distance(data: dict) -> Optional[int]:
        """Read the actual distance NIC used back from the response."""
        for key in ("distance", "Distance", "transDistance", "TransDistance", "actualDist"):
            val = data.get(key)
            if val:
                try:
                    return int(float(val))
                except (ValueError, TypeError):
                    continue
        return None

    @staticmethod
    def _calc_validity(distance_km: int, km_per_day: int = 200) -> datetime:
        """NIC rule: 1 day validity per `km_per_day` km (or part thereof), min 1 day, ending 23:59."""
        import math
        days = max(1, math.ceil((distance_km or 0) / (km_per_day or 200)))
        return (datetime.now() + timedelta(days=days)).replace(
            hour=23, minute=59, second=0, microsecond=0
        )

    @staticmethod
    def _parse_validity(value) -> Optional[datetime]:
        if not value:
            return None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y %I:%M:%S %p",
                    "%d/%m/%Y %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(str(value), fmt)
            except (ValueError, TypeError):
                continue
        return None

    @staticmethod
    async def generate(db: Session, payload: EWayBillGenerateRequest, user_id: int) -> dict:
        invoice = db.query(Invoice).filter(Invoice.id == payload.invoice_id).first()
        if not invoice:
            raise HTTPException(status_code=404, detail="Invoice not found")
        if invoice.is_cancelled:
            raise HTTPException(status_code=400, detail="Cannot generate E-Way Bill for cancelled invoice")
        eway_threshold = float(ConfigService.get_decimal(
            db, "gst.eway_bill_threshold",
            as_of=invoice.invoice_date, default=settings.EWAY_BILL_THRESHOLD,
        ))
        if float(invoice.total_amount) < eway_threshold:
            raise HTTPException(
                status_code=400,
                detail=f"Invoice value ₹{invoice.total_amount} is below E-Way Bill threshold ₹{eway_threshold}"
            )
        eway_km_per_day = ConfigService.get_int(
            db, "gst.eway_km_per_day",
            as_of=invoice.invoice_date, default=200,
        )

        ship_addr = db.query(CustomerAddress).filter(CustomerAddress.id == invoice.shipping_address_id).first()
        # Option A: when no manual distance, send 0 so NIC auto-calculates the
        # actual pin-to-pin road distance and returns it. The offline estimate
        # is only a fallback for mock mode / responses that omit the distance.
        manual_distance = payload.distance_km
        sent_distance = manual_distance if manual_distance else 0
        company = CompanyService.identity(db)
        estimated_distance = EWayBillService._calc_distance(
            company.pincode,
            ship_addr.pincode if ship_addr else "000000"
        )

        vehicle_number = payload.vehicle_number
        transporter_name = payload.transporter_name
        if payload.vehicle_id:
            veh = db.query(Vehicle).filter(Vehicle.id == payload.vehicle_id).first()
            if veh:
                vehicle_number = veh.vehicle_number
        if payload.transporter_id:
            trans = db.query(Transporter).filter(Transporter.id == payload.transporter_id).first()
            if trans:
                transporter_name = trans.name

        eway_payload = {
            "supplyType": "O",
            "subSupplyType": "1",
            "docType": "INV",
            "docNo": invoice.invoice_number,
            "docDate": invoice.invoice_date.strftime("%d/%m/%Y"),
            "fromGstin": company.gstin,
            "fromTrdName": company.name,
            "fromAddr1": company.address,
            "fromPlace": company.city,
            "fromPincode": int(company.pincode),
            "fromStateCode": company.state_code,
            "toGstin": (db.query(Customer).filter(Customer.id == invoice.customer_id).first().gstin or "URP"),
            "toTrdName": (db.query(Customer).filter(Customer.id == invoice.customer_id).first().trade_name),
            "toAddr1": ship_addr.address_line1 if ship_addr else "",
            "toPlace": ship_addr.city if ship_addr else "",
            "toPincode": int(ship_addr.pincode) if ship_addr and ship_addr.pincode else 0,
            "toStateCode": ship_addr.state_code if ship_addr and ship_addr.state_code else company.state_code,
            "totalValue": float(invoice.total_amount),
            "cgstValue": float(invoice.total_cgst),
            "sgstValue": float(invoice.total_sgst),
            "igstValue": float(invoice.total_igst),
            "transporterName": transporter_name or "",
            "transDistance": sent_distance,
            "vehicleNo": vehicle_number or "",
            "vehicleType": "R",
            "transMode": {"road": "1", "rail": "2", "air": "3", "ship": "4"}.get(
                payload.transport_mode, "1"
            ),
        }

        eway_bill_number = None
        valid_upto = None
        error_message = None
        response_data = None
        success = False
        # Distance actually recorded on the bill: manual value, else NIC's
        # auto-calculated value from the response, else the offline estimate.
        final_distance = manual_distance or estimated_distance

        from_pincode = str(company.pincode)
        to_pincode = ship_addr.pincode if ship_addr else None

        if EInvoiceService._should_mock():
            eway_bill_number = f"EWB{payload.invoice_id:012d}"
            valid_upto = EWayBillService._calc_validity(final_distance, eway_km_per_day)
            response_data = {"mock": True, "transDistance": final_distance}
            success = True
        else:
            try:
                async with httpx.AsyncClient(timeout=30) as client:
                    resp = await client.post(
                        f"{settings.CLEARTAX_API_URL}/v2/eWaybill/generate",
                        json=eway_payload,
                        headers={"x-cleartax-auth-token": settings.CLEARTAX_AUTH_TOKEN}
                    )
                    if resp.status_code == 200:
                        data = resp.json()
                        response_data = data
                        eway_bill_number = data.get("ewbNo") or data.get("EwbNo")
                        nic_distance = EWayBillService._extract_distance(data)
                        if nic_distance:
                            final_distance = nic_distance
                        valid_upto = (
                            EWayBillService._parse_validity(data.get("validUpto") or data.get("ValidUpto"))
                            or EWayBillService._calc_validity(final_distance, eway_km_per_day)
                        )
                        success = True
                    else:
                        try:
                            err_data = resp.json()
                            error_message = err_data.get("message") or err_data.get("error") or resp.text
                        except Exception:
                            error_message = resp.text or f"API error: HTTP {resp.status_code}"
            except Exception as e:
                error_message = str(e) or "Network error connecting to Cleartax API"

        log = EWayBillLog(
            invoice_id=payload.invoice_id,
            eway_bill_number=str(eway_bill_number) if eway_bill_number else None,
            irn=invoice.irn,
            vehicle_number=vehicle_number,
            transporter_name=transporter_name,
            transport_mode=payload.transport_mode,
            distance_km=final_distance,
            from_pincode=from_pincode,
            to_pincode=to_pincode,
            valid_upto=valid_upto,
            status=EWayBillStatus.generated if success else EWayBillStatus.failed,
            request_payload=json.dumps(eway_payload),
            response_payload=json.dumps(response_data) if response_data is not None else None,
            error_message=error_message,
            created_by=user_id,
        )
        db.add(log)
        if success and eway_bill_number:
            invoice.eway_bill_number = str(eway_bill_number)
            invoice.eway_bill_status = EWayBillStatus.generated
        else:
            invoice.eway_bill_status = EWayBillStatus.failed
        db.commit()

        audit(db, user_id, "create", "gst",
              f"Generated E-Way Bill {eway_bill_number} for invoice {invoice.invoice_number}" if success
              else f"E-Way Bill generation failed for invoice {invoice.invoice_number}: {error_message}",
              record_type="eway_bill", record_id=payload.invoice_id)

        return {
            "invoice_id": payload.invoice_id,
            "invoice_number": invoice.invoice_number,
            "eway_bill_number": eway_bill_number,
            "valid_upto": valid_upto.strftime("%Y-%m-%d %H:%M:%S") if valid_upto else None,
            "error_message": (error_message or "Unknown error") if not success else None,
        }

    @staticmethod
    def list_eway_logs(db: Session, page: int = 1, page_size: int = 20) -> dict:
        q = db.query(EWayBillLog).order_by(desc(EWayBillLog.created_at))
        result = paginate(q, page, page_size)
        items = []
        for log in result["items"]:
            inv = db.query(Invoice).filter(Invoice.id == log.invoice_id).first()
            items.append({
                "id": log.id,
                "invoice_id": log.invoice_id,
                "invoice_number": inv.invoice_number if inv else None,
                "eway_bill_number": log.eway_bill_number,
                "vehicle_number": log.vehicle_number,
                "transporter_name": log.transporter_name,
                "distance_km": log.distance_km,
                "created_at": log.created_at,
            })
        result["items"] = items
        return result


class GSTR1Service:

    @staticmethod
    def generate(db: Session, period: str, financial_year: str) -> dict:
        """Generate GSTR-1 data for a given period (MM-YYYY)."""
        try:
            month, year = period.split("-")
            month_int = int(month)
            year_int = int(year)
            from_date = date(year_int, month_int, 1)
            if month_int == 12:
                to_date = date(year_int + 1, 1, 1)
            else:
                to_date = date(year_int, month_int + 1, 1)
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid period format. Use MM-YYYY e.g. 01-2026")

        invoices = db.query(Invoice).filter(
            Invoice.invoice_date >= from_date,
            Invoice.invoice_date < to_date,
            Invoice.is_cancelled == False,
            Invoice.document_type.in_([
                DocumentType.b2b_invoice, DocumentType.b2c_invoice,
                DocumentType.credit_note,
            ]),
        ).all()

        b2cl_threshold = float(ConfigService.get_decimal(
            db, "gst.gstr1_b2cl_threshold",
            as_of=from_date, default=settings.GSTR1_B2CL_THRESHOLD,
        ))
        company_state_code = CompanyService.identity(db).state_code

        b2b, b2c, credit_notes = [], [], []
        b2cl, b2cs_map = [], {}
        total_taxable = total_igst = total_cgst = total_sgst = Decimal("0")
        hsn_map = {}

        for inv in invoices:
            customer = db.query(Customer).filter(Customer.id == inv.customer_id).first()
            is_cn = inv.document_type == DocumentType.credit_note
            # Credit notes are stored with positive amounts but reduce outward
            # liability — net them off totals/HSN and report them under 9B.
            sign = Decimal("-1") if is_cn else Decimal("1")

            if is_cn:
                if customer and customer.gstin:
                    original = db.query(Invoice).filter(
                        Invoice.id == inv.original_invoice_id
                    ).first() if inv.original_invoice_id else None
                    credit_notes.append({
                        "note_number": inv.invoice_number,
                        "note_date": inv.invoice_date,
                        "note_type": "C",
                        "original_invoice_number": original.invoice_number if original else None,
                        "original_invoice_date": original.invoice_date if original else None,
                        "customer_gstin": customer.gstin,
                        "customer_name": customer.trade_name,
                        "note_value": inv.total_amount,
                        "taxable_value": inv.taxable_amount,
                        "igst": inv.total_igst,
                        "cgst": inv.total_cgst,
                        "sgst": inv.total_sgst,
                        "registered": True,
                    })
                else:
                    # Unregistered B2C credit note → netted into B2CS (Table 7),
                    # not reported as a separate CDNUR row.
                    ship_addr = db.query(CustomerAddress).filter(
                        CustomerAddress.id == inv.shipping_address_id
                    ).first()
                    pos_state = ship_addr.state if ship_addr else None
                    pos_code = ship_addr.state_code if ship_addr else None
                    is_inter = pos_code is not None and pos_code != company_state_code
                    for item in inv.items:
                        rate = item.gst_percent if item.gst_percent is not None else Decimal("0")
                        skey = (pos_code, str(rate), "inter" if is_inter else "intra")
                        if skey not in b2cs_map:
                            b2cs_map[skey] = {
                                "state": pos_state,
                                "state_code": pos_code,
                                "gst_rate": rate,
                                "supply_type": "inter" if is_inter else "intra",
                                "taxable_value": Decimal("0"),
                                "igst": Decimal("0"),
                                "cgst": Decimal("0"),
                                "sgst": Decimal("0"),
                            }
                        b2cs_map[skey]["taxable_value"] -= item.taxable_amount
                        b2cs_map[skey]["igst"] -= item.igst_amount
                        b2cs_map[skey]["cgst"] -= item.cgst_amount
                        b2cs_map[skey]["sgst"] -= item.sgst_amount
            else:
                row = {
                    "invoice_number": inv.invoice_number,
                    "invoice_date": inv.invoice_date,
                    "customer_gstin": customer.gstin if customer else None,
                    "customer_name": customer.trade_name if customer else None,
                    "invoice_value": inv.total_amount,
                    "taxable_value": inv.taxable_amount,
                    "igst": inv.total_igst,
                    "cgst": inv.total_cgst,
                    "sgst": inv.total_sgst,
                    "irn": inv.irn,
                }
                if inv.document_type == DocumentType.b2b_invoice and customer and customer.gstin:
                    b2b.append(row)
                else:
                    ship_addr = db.query(CustomerAddress).filter(
                        CustomerAddress.id == inv.shipping_address_id
                    ).first()
                    pos_state = ship_addr.state if ship_addr else None
                    pos_code = ship_addr.state_code if ship_addr else None
                    row["state"] = pos_state
                    b2c.append(row)

                    # B2CL = inter-state and above threshold (Table 5, invoice-wise);
                    # everything else rolls into B2CS (Table 7, by POS + rate).
                    is_inter = pos_code is not None and pos_code != company_state_code
                    if is_inter and float(inv.total_amount or 0) > b2cl_threshold:
                        b2cl.append({
                            "invoice_number": inv.invoice_number,
                            "invoice_date": inv.invoice_date,
                            "state": pos_state,
                            "state_code": pos_code,
                            "invoice_value": inv.total_amount,
                            "taxable_value": inv.taxable_amount,
                            "igst": inv.total_igst,
                            "cgst": inv.total_cgst,
                            "sgst": inv.total_sgst,
                        })
                    else:
                        for item in inv.items:
                            rate = item.gst_percent if item.gst_percent is not None else Decimal("0")
                            skey = (pos_code, str(rate), "inter" if is_inter else "intra")
                            if skey not in b2cs_map:
                                b2cs_map[skey] = {
                                    "state": pos_state,
                                    "state_code": pos_code,
                                    "gst_rate": rate,
                                    "supply_type": "inter" if is_inter else "intra",
                                    "taxable_value": Decimal("0"),
                                    "igst": Decimal("0"),
                                    "cgst": Decimal("0"),
                                    "sgst": Decimal("0"),
                                }
                            b2cs_map[skey]["taxable_value"] += item.taxable_amount
                            b2cs_map[skey]["igst"] += item.igst_amount
                            b2cs_map[skey]["cgst"] += item.cgst_amount
                            b2cs_map[skey]["sgst"] += item.sgst_amount

            total_taxable += sign * inv.taxable_amount
            total_igst += sign * inv.total_igst
            total_cgst += sign * inv.total_cgst
            total_sgst += sign * inv.total_sgst

            # HSN summary (net of credit notes). Key on (HSN, rate, UQC) so the
            # same HSN sold at different GST rates is not merged into one row.
            for item in inv.items:
                prod = db.query(Product).filter(Product.id == item.product_id).first()
                uom = prod.unit_of_measure if prod else "Nos"
                rate = item.gst_percent if item.gst_percent is not None else Decimal("0")
                key = (item.hsn_code, str(rate), uom)
                if key not in hsn_map:
                    hsn_map[key] = {
                        "hsn_code": item.hsn_code,
                        "description": prod.part_name if prod else None,
                        "uom": uom,
                        "total_quantity": Decimal("0"),
                        "taxable_value": Decimal("0"),
                        "igst": Decimal("0"),
                        "cgst": Decimal("0"),
                        "sgst": Decimal("0"),
                        "gst_rate": rate,
                    }
                hsn_map[key]["total_quantity"] += sign * item.quantity
                hsn_map[key]["taxable_value"] += sign * item.taxable_amount
                hsn_map[key]["igst"] += sign * item.igst_amount
                hsn_map[key]["cgst"] += sign * item.cgst_amount
                hsn_map[key]["sgst"] += sign * item.sgst_amount

        b2cs_summary = [
            {**v,
             "taxable_value": v["taxable_value"].quantize(Decimal("0.01")),
             "igst": v["igst"].quantize(Decimal("0.01")),
             "cgst": v["cgst"].quantize(Decimal("0.01")),
             "sgst": v["sgst"].quantize(Decimal("0.01"))}
            for v in b2cs_map.values()
        ]

        return {
            "period": period,
            "financial_year": financial_year,
            "b2b_invoices": b2b,
            "b2c_invoices": b2c,
            "b2cl_invoices": b2cl,
            "b2cs_summary": b2cs_summary,
            "hsn_summary": list(hsn_map.values()),
            "credit_notes": credit_notes,
            "total_taxable": total_taxable.quantize(Decimal("0.01")),
            "total_igst": total_igst.quantize(Decimal("0.01")),
            "total_cgst": total_cgst.quantize(Decimal("0.01")),
            "total_sgst": total_sgst.quantize(Decimal("0.01")),
            "total_tax": (total_igst + total_cgst + total_sgst).quantize(Decimal("0.01")),
            "invoice_count": len(b2b) + len(b2c),
            "b2b_count": len(b2b),
            "b2c_count": len(b2c),
            "b2cl_count": len(b2cl),
            "b2cs_count": len(b2cs_summary),
            "b2cl_threshold": settings.GSTR1_B2CL_THRESHOLD,
            "credit_note_count": len(credit_notes),
        }


class GSTR1ExportService:
    """GSTR-1 downloads: human-readable Excel and GSTN offline-tool JSON.

    The JSON aims to match the GST offline utility schema. Assumptions:
    cess is 0 (no cess tracked), all invoices are regular ('R') with no
    reverse charge, and place-of-supply is the shipping-address state code.
    Validate the `version` string against your current offline tool before
    filing live.
    """

    @staticmethod
    def _period_bounds(period: str):
        try:
            month, year = period.split("-")
            m, y = int(month), int(year)
            from_date = date(y, m, 1)
            to_date = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
            return from_date, to_date
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid period. Use MM-YYYY")

    @staticmethod
    def _fmt_dt(d) -> Optional[str]:
        return d.strftime("%d-%m-%Y") if d else None

    @staticmethod
    def _itms_by_rate(inv) -> List[dict]:
        groups = {}
        for it in inv.items:
            rt = float(it.gst_percent or 0)
            g = groups.setdefault(rt, {"rt": rt, "txval": 0.0, "iamt": 0.0,
                                       "camt": 0.0, "samt": 0.0, "csamt": 0.0})
            g["txval"] += float(it.taxable_amount or 0)
            g["iamt"] += float(it.igst_amount or 0)
            g["camt"] += float(it.cgst_amount or 0)
            g["samt"] += float(it.sgst_amount or 0)
        return [
            {"num": i, "itm_det": {k: round(v, 2) if isinstance(v, float) else v
                                   for k, v in g.items()}}
            for i, g in enumerate(groups.values(), 1)
        ]

    @staticmethod
    def build_excel(db: Session, period: str, financial_year: str) -> bytes:
        import io as _io
        from openpyxl import Workbook

        data = GSTR1Service.generate(db, period, financial_year)
        wb = Workbook()

        ws = wb.active
        ws.title = "Summary"
        ws.append(["GSTR-1", f"Period {period}", f"FY {financial_year}"])
        ws.append([])
        for label, val in [
            ("Total Invoices", data["invoice_count"]),
            ("Taxable Value", float(data["total_taxable"])),
            ("IGST", float(data["total_igst"])),
            ("CGST", float(data["total_cgst"])),
            ("SGST", float(data["total_sgst"])),
            ("Total Tax", float(data["total_tax"])),
            ("B2B Count", data["b2b_count"]),
            ("B2CL Count", data["b2cl_count"]),
            ("B2CS Count", data["b2cs_count"]),
            ("Credit Notes", data["credit_note_count"]),
        ]:
            ws.append([label, val])

        ws = wb.create_sheet("B2B")
        ws.append(["Invoice No", "Date", "Customer GSTIN", "Customer",
                   "Invoice Value", "Taxable", "IGST", "CGST", "SGST", "IRN"])
        for r in data["b2b_invoices"]:
            ws.append([r["invoice_number"], str(r["invoice_date"]), r["customer_gstin"],
                       r["customer_name"], float(r["invoice_value"]), float(r["taxable_value"]),
                       float(r["igst"]), float(r["cgst"]), float(r["sgst"]), r.get("irn")])

        ws = wb.create_sheet("B2CL")
        ws.append(["Invoice No", "Date", "POS State", "State Code",
                   "Invoice Value", "Taxable", "IGST"])
        for r in data["b2cl_invoices"]:
            ws.append([r["invoice_number"], str(r["invoice_date"]), r["state"],
                       r["state_code"], float(r["invoice_value"]),
                       float(r["taxable_value"]), float(r["igst"])])

        ws = wb.create_sheet("B2CS")
        ws.append(["POS State", "State Code", "Supply Type", "Rate %",
                   "Taxable", "IGST", "CGST", "SGST"])
        for r in data["b2cs_summary"]:
            ws.append([r["state"], r["state_code"], r["supply_type"], float(r["gst_rate"]),
                       float(r["taxable_value"]), float(r["igst"]),
                       float(r["cgst"]), float(r["sgst"])])

        ws = wb.create_sheet("HSN")
        ws.append(["HSN", "Description", "UOM", "Quantity", "Rate %",
                   "Taxable", "IGST", "CGST", "SGST"])
        for r in data["hsn_summary"]:
            ws.append([r["hsn_code"], r["description"], r["uom"], float(r["total_quantity"]),
                       float(r["gst_rate"]), float(r["taxable_value"]), float(r["igst"]),
                       float(r["cgst"]), float(r["sgst"])])

        ws = wb.create_sheet("Credit Notes")
        ws.append(["Note No", "Date", "Against Invoice", "Customer GSTIN", "Customer",
                   "Type", "Note Value", "Taxable", "IGST", "CGST", "SGST"])
        for r in data["credit_notes"]:
            ws.append([r["note_number"], str(r["note_date"]), r.get("original_invoice_number"),
                       r.get("customer_gstin"), r.get("customer_name"),
                       "CDNR" if r.get("registered") else "CDNUR", float(r["note_value"]),
                       float(r["taxable_value"]), float(r["igst"]),
                       float(r["cgst"]), float(r["sgst"])])

        buf = _io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @staticmethod
    def build_gov_json(db: Session, period: str, financial_year: str) -> dict:
        from_date, to_date = GSTR1ExportService._period_bounds(period)
        invoices = db.query(Invoice).filter(
            Invoice.invoice_date >= from_date,
            Invoice.invoice_date < to_date,
            Invoice.is_cancelled == False,
            Invoice.document_type.in_([
                DocumentType.b2b_invoice, DocumentType.b2c_invoice,
                DocumentType.credit_note,
            ]),
        ).all()

        company = CompanyService.identity(db)
        company_pos = str(company.state_code)
        b2cl_threshold = float(ConfigService.get_decimal(
            db, "gst.gstr1_b2cl_threshold",
            as_of=from_date, default=settings.GSTR1_B2CL_THRESHOLD,
        ))
        gstr1_version = ConfigService.get_string(
            db, "gst.gstr1_json_version", as_of=from_date, default="GST3.1.4",
        )
        uqc_map = GstMasterService.uqc_map(db, from_date)
        b2b_by_ctin, b2cl_by_pos, b2cs_agg, cdnr_by_ctin, hsn_agg = {}, {}, {}, {}, {}

        for inv in invoices:
            customer = db.query(Customer).filter(Customer.id == inv.customer_id).first()
            ship = db.query(CustomerAddress).filter(
                CustomerAddress.id == inv.shipping_address_id).first()
            pos = (str(inv.place_of_supply) if getattr(inv, "place_of_supply", None)
                   else (str(ship.state_code) if ship and ship.state_code else company_pos))
            is_cn = inv.document_type == DocumentType.credit_note
            sign = -1.0 if is_cn else 1.0

            if is_cn:
                if customer and customer.gstin:
                    cdnr_by_ctin.setdefault(customer.gstin, []).append({
                        "ntty": "C",
                        "nt_num": inv.invoice_number,
                        "nt_dt": GSTR1ExportService._fmt_dt(inv.invoice_date),
                        "val": float(inv.total_amount or 0),
                        "rchrg": "N",
                        "inv_typ": "R",
                        "itms": GSTR1ExportService._itms_by_rate(inv),
                    })
                else:
                    # Unregistered B2C credit note → netted into B2CS (Table 7).
                    is_inter = pos != company_pos
                    for it in inv.items:
                        rt = float(it.gst_percent or 0)
                        sply = "INTER" if is_inter else "INTRA"
                        a = b2cs_agg.setdefault((sply, pos, rt), {
                            "sply_ty": sply, "pos": pos, "typ": "OE", "rt": rt,
                            "txval": 0.0, "iamt": 0.0, "camt": 0.0, "samt": 0.0, "csamt": 0.0,
                        })
                        a["txval"] -= float(it.taxable_amount or 0)
                        a["iamt"] -= float(it.igst_amount or 0)
                        a["camt"] -= float(it.cgst_amount or 0)
                        a["samt"] -= float(it.sgst_amount or 0)
            elif inv.document_type == DocumentType.b2b_invoice and customer and customer.gstin:
                b2b_by_ctin.setdefault(customer.gstin, []).append({
                    "inum": inv.invoice_number,
                    "idt": GSTR1ExportService._fmt_dt(inv.invoice_date),
                    "val": float(inv.total_amount or 0),
                    "pos": pos,
                    "rchrg": "N",
                    "inv_typ": "R",
                    "itms": GSTR1ExportService._itms_by_rate(inv),
                })
            else:
                is_inter = pos != company_pos
                if is_inter and float(inv.total_amount or 0) > b2cl_threshold:
                    b2cl_by_pos.setdefault(pos, []).append({
                        "inum": inv.invoice_number,
                        "idt": GSTR1ExportService._fmt_dt(inv.invoice_date),
                        "val": float(inv.total_amount or 0),
                        "itms": GSTR1ExportService._itms_by_rate(inv),
                    })
                else:
                    for it in inv.items:
                        rt = float(it.gst_percent or 0)
                        sply = "INTER" if is_inter else "INTRA"
                        a = b2cs_agg.setdefault((sply, pos, rt), {
                            "sply_ty": sply, "pos": pos, "typ": "OE", "rt": rt,
                            "txval": 0.0, "iamt": 0.0, "camt": 0.0, "samt": 0.0, "csamt": 0.0,
                        })
                        a["txval"] += float(it.taxable_amount or 0)
                        a["iamt"] += float(it.igst_amount or 0)
                        a["camt"] += float(it.cgst_amount or 0)
                        a["samt"] += float(it.sgst_amount or 0)

            for it in inv.items:
                prod = db.query(Product).filter(Product.id == it.product_id).first()
                uqc = normalize_uqc(prod.unit_of_measure if prod else None, uqc_map)
                rt = float(it.gst_percent or 0)
                a = hsn_agg.setdefault((it.hsn_code, rt, uqc), {
                    "hsn_sc": it.hsn_code, "desc": (prod.part_name if prod else None),
                    "uqc": uqc, "rt": rt, "qty": 0.0, "txval": 0.0,
                    "iamt": 0.0, "camt": 0.0, "samt": 0.0, "csamt": 0.0,
                })
                a["qty"] += sign * float(it.quantity or 0)
                a["txval"] += sign * float(it.taxable_amount or 0)
                a["iamt"] += sign * float(it.igst_amount or 0)
                a["camt"] += sign * float(it.cgst_amount or 0)
                a["samt"] += sign * float(it.sgst_amount or 0)

        def _round(d):
            return {k: (round(v, 2) if isinstance(v, float) else v) for k, v in d.items()}

        hsn_data = [{"num": i, **_round(a)} for i, a in enumerate(hsn_agg.values(), 1)]

        return {
            "gstin": company.gstin,
            "fp": period.replace("-", ""),
            "gt": 0,
            "cur_gt": 0,
            "version": gstr1_version,
            "hash": "hash",
            "b2b": [{"ctin": c, "inv": v} for c, v in b2b_by_ctin.items()],
            "b2cl": [{"pos": p, "inv": v} for p, v in b2cl_by_pos.items()],
            "b2cs": [_round(a) for a in b2cs_agg.values()],
            "cdnr": [{"ctin": c, "nt": n} for c, n in cdnr_by_ctin.items()],
            "hsn": {"data": hsn_data},
        }


class GSTR2BService:

    @staticmethod
    def reconcile(db: Session, payload: GSTR2BImport) -> dict:
        matched, unmatched_gstr2b, unmatched_books = [], [], []
        itc_claimed = Decimal("0")
        itc_at_risk = Decimal("0")

        for entry in payload.entries:
            purchase = db.query(Purchase).filter(
                Purchase.vendor_invoice_number == entry.invoice_number,
                Purchase.is_cancelled == False,
            ).first()

            if purchase:
                vendor = db.query(Vendor).filter(Vendor.id == purchase.vendor_id).first()
                if vendor and vendor.gstin == entry.supplier_gstin:
                    matched.append({
                        "supplier_gstin": entry.supplier_gstin,
                        "supplier_name": entry.supplier_name,
                        "invoice_number": entry.invoice_number,
                        "invoice_date": entry.invoice_date,
                        "invoice_value": entry.invoice_value,
                        "igst": entry.igst,
                        "cgst": entry.cgst,
                        "sgst": entry.sgst,
                        "purchase_id": purchase.id,
                    })
                    itc_claimed += entry.igst + entry.cgst + entry.sgst
                else:
                    unmatched_gstr2b.append({**entry.__dict__, "reason": "GSTIN mismatch"})
                    itc_at_risk += entry.igst + entry.cgst + entry.sgst
            else:
                unmatched_gstr2b.append({
                    "supplier_gstin": entry.supplier_gstin,
                    "supplier_name": entry.supplier_name,
                    "invoice_number": entry.invoice_number,
                    "invoice_date": str(entry.invoice_date),
                    "invoice_value": float(entry.invoice_value),
                    "reason": "Not found in books",
                })
                itc_at_risk += entry.igst + entry.cgst + entry.sgst

        # Find purchases not in GSTR-2B
        gstr2b_inv_numbers = {e.invoice_number for e in payload.entries}
        unmatched_purchases = db.query(Purchase).filter(
            Purchase.is_cancelled == False,
            ~Purchase.vendor_invoice_number.in_(gstr2b_inv_numbers)
        ).limit(100).all()
        for p in unmatched_purchases:
            v = db.query(Vendor).filter(Vendor.id == p.vendor_id).first()
            unmatched_books.append({
                "purchase_id": p.id,
                "vendor_name": v.trade_name if v else None,
                "invoice_number": p.vendor_invoice_number,
                "invoice_date": str(p.invoice_date),
                "total_amount": float(p.total_amount),
                "reason": "Not in GSTR-2B",
            })

        return {
            "period": payload.period,
            "financial_year": payload.financial_year,
            "matched": matched,
            "unmatched_in_gstr2b": unmatched_gstr2b,
            "unmatched_in_books": unmatched_books[:50],
            "matched_count": len(matched),
            "unmatched_gstr2b_count": len(unmatched_gstr2b),
            "unmatched_books_count": len(unmatched_books),
            "total_itc_available": (itc_claimed + itc_at_risk).quantize(Decimal("0.01")),
            "total_itc_claimed": itc_claimed.quantize(Decimal("0.01")),
            "itc_at_risk": itc_at_risk.quantize(Decimal("0.01")),
        }


class GSTR3BService:

    @staticmethod
    def generate(db: Session, period: str, financial_year: str) -> dict:
        try:
            month, year = period.split("-")
            month_int = int(month)
            year_int = int(year)
            from_date = date(year_int, month_int, 1)
            if month_int == 12:
                to_date = date(year_int + 1, 1, 1)
            else:
                to_date = date(year_int, month_int + 1, 1)
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid period. Use MM-YYYY")

        # Output tax from sales invoices
        sales = db.query(Invoice).filter(
            Invoice.invoice_date >= from_date,
            Invoice.invoice_date < to_date,
            Invoice.is_cancelled == False,
            Invoice.document_type.in_([DocumentType.b2b_invoice, DocumentType.b2c_invoice]),
        ).all()

        out_igst = out_cgst = out_sgst = Decimal("0")
        for inv in sales:
            out_igst += inv.total_igst
            out_cgst += inv.total_cgst
            out_sgst += inv.total_sgst

        # ITC from purchases
        purchases = db.query(Purchase).filter(
            Purchase.invoice_date >= from_date,
            Purchase.invoice_date < to_date,
            Purchase.is_cancelled == False,
        ).all()

        itc_igst = itc_cgst = itc_sgst = Decimal("0")
        for p in purchases:
            itc_igst += p.total_igst
            itc_cgst += p.total_cgst
            itc_sgst += p.total_sgst

        total_output = out_igst + out_cgst + out_sgst
        total_itc = itc_igst + itc_cgst + itc_sgst
        net_payable = max(Decimal("0"), total_output - total_itc)
        net_igst = max(Decimal("0"), out_igst - itc_igst)
        net_cgst = max(Decimal("0"), out_cgst - itc_cgst)
        net_sgst = max(Decimal("0"), out_sgst - itc_sgst)

        return {
            "period": period,
            "financial_year": financial_year,
            "output_igst": out_igst.quantize(Decimal("0.01")),
            "output_cgst": out_cgst.quantize(Decimal("0.01")),
            "output_sgst": out_sgst.quantize(Decimal("0.01")),
            "total_output_tax": total_output.quantize(Decimal("0.01")),
            "itc_igst": itc_igst.quantize(Decimal("0.01")),
            "itc_cgst": itc_cgst.quantize(Decimal("0.01")),
            "itc_sgst": itc_sgst.quantize(Decimal("0.01")),
            "total_itc": total_itc.quantize(Decimal("0.01")),
            "net_tax_payable": net_payable.quantize(Decimal("0.01")),
            "net_igst_payable": net_igst.quantize(Decimal("0.01")),
            "net_cgst_payable": net_cgst.quantize(Decimal("0.01")),
            "net_sgst_payable": net_sgst.quantize(Decimal("0.01")),
            "invoice_count": len(sales),
            "purchase_count": len(purchases),
        }